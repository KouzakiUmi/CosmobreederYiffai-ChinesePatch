"""Review workbook, Grok batch preparation/merge, validation and safe text rendering."""
from pathlib import Path
import argparse
import collections
import datetime
import json
import re
import sys
from projectlib import *


def prepare():
    units=load_units();occ=load_occurrences()
    if not (PROJECT/'translations/zh-CN.jsonl').exists():
        write_jsonl(PROJECT/'translations/zh-CN.jsonl',[{'id':u['id'],'source_sha256':u['source_sha256'],'draft':'','translation':'',
            'status':'untranslated','model':'','notes':'','reviewer':'','usage_approved':False} for u in units])
    if not (PROJECT/'review/occurrence_overrides.csv').exists():
        by_id={u['id']:u for u in units}
        write_csv(PROJECT/'review/occurrence_overrides.csv',
            [{'occurrence_id':o['occurrence_id'],'unit_id':o['unit_id'],'location':o['location'],'source':by_id[o['unit_id']]['source'],
              'category':o['category'],'variant':o['variant'],'speaker':o['speaker'],'translation_override':'','status':'untranslated','notes':''} for o in occ],
            ['occurrence_id','unit_id','location','source','category','variant','speaker','translation_override','status','notes'])
    if not (PROJECT/'review/glossary.csv').exists():
        old=read_csv(PROJECT/'reference/glossary-seed.csv')
        write_csv(PROJECT/'review/glossary.csv',[{**r,'status':'proposed'} for r in old],['term','usage','proposed_zh','note','status'])
    if (PROJECT/'batches/manifest.json').exists():raise SystemExit('Batch manifest exists; do not regenerate assigned batches.')
    order={'ui':0,'vsn':1,'dialogue':2,'code':3,'image':4}
    # Scene/entry context is retained; no cross-protagonist or cross-speaker deduplication.
    units=sorted(units,key=lambda u:(order[u['category']],u['context_examples'][0] if u['context_examples'] else '',u['id']))
    batches=[];current=[];chars=0
    def save():
        if not current:return
        bid='batch_%03d'%len(batches)
        write_jsonl(PROJECT/'batches'/(bid+'.jsonl'),current)
        batches.append({'batch_id':bid,'path':'batches/'+bid+'.jsonl','draft_path':'drafts/'+bid+'.jsonl',
                        'count':len(current),'source_characters':sum(len(r['source']) for r in current),
                        'ids':[r['id'] for r in current]})
    for u in units:
        if current and (len(current)>=120 or chars+len(u['source'])>14000):save();current=[];chars=0
        current.append({k:u[k] for k in ['id','category','source','source_sha256','variant','speaker','context_examples','protected_tokens','markers','width_hint_px','review_flags']});chars+=len(u['source'])
    save();write_json(PROJECT/'batches/manifest.json',{'model':read_json(PROJECT/'project.json')['translation_route'],'batches':batches})
    print('Prepared',len(units),'translation units in',len(batches),'bounded batches')


def merged_value(record,include_drafts=False):
    if record.get('status') in ('reviewed','approved') and record.get('translation'):return record['translation']
    if include_drafts and record.get('draft'):return record['draft']
    return None


def merge_drafts():
    units={u['id']:u for u in load_units()};records=current_translations();manifest=read_json(PROJECT/'batches/manifest.json')
    accepted=[];pending=[];partial=[];accepted_parts=[];errors=[]
    for b in manifest['batches']:
        full=PROJECT/b['draft_path'];sources=[];covered=0
        if full.exists() and (not b.get('require_ready') or full.with_suffix('.ready.json').exists()):sources=[(full,b['ids'])]
        else:
            for start in range(0,b['count'],40):
                p=PROJECT/'drafts/parts'/(b['batch_id']+f'_{start:03d}.jsonl');ready=p.with_suffix('.ready.json')
                if ready.exists():sources.append((p,b['ids'][start:start+40]))
        for p,expected_ids in sources:
            try:
                origin_model=draft_model_label(p)
                if p!=full or b.get('require_ready'):
                    seal=read_json(p.with_suffix('.ready.json'))
                    if seal['draft_sha256']!=digest(p.read_bytes()) or seal['count']!=len(expected_ids):raise ValueError('Committed part fingerprint mismatch')
                    if seal.get('origin_model',origin_model)!=origin_model:raise ValueError('Committed part model provenance mismatch')
                rows=read_jsonl(p);seen=[r['id'] for r in rows]
                if seen!=expected_ids:raise ValueError('IDs missing, duplicated, reordered or outside assigned input')
                for row in rows:
                    u=units[row['id']]
                    if row.get('source_sha256')!=u['source_sha256']:raise ValueError('source hash mismatch '+u['id'])
                    if row.get('status')=='blocked':
                        if not row.get('notes') or row.get('translation','')!='':raise ValueError('Blocked row needs explicit reason '+u['id'])
                    elif not isinstance(row.get('translation'),str) or not row['translation'].strip():raise ValueError('empty translation '+u['id'])
                    if not isinstance(row.get('notes',''),str):raise ValueError('notes must be a string')
                    if records[row['id']].get('draft') and records[row['id']]['draft']!=row['translation']:raise ValueError('Existing draft differs '+row['id'])
                for row in rows:
                    rec=records[row['id']];human_work=bool(rec.get('translation') or rec.get('reviewer') or rec['status'] in ('reviewed','approved'))
                    rec['draft']=row['translation'];rec['model']=origin_model;rec['draft_notes']=row.get('notes','')
                    if not human_work:
                        if row.get('status')=='blocked':rec['status']='blocked'
                        elif rec['status'] in ('untranslated','blocked'):rec['status']='draft'
                        rec['notes']=row.get('notes','')
                    rec['draft_batch']=b['batch_id'];rec['draft_source_path']=p.relative_to(PROJECT).as_posix()
                covered+=len(rows)
                if p!=full:accepted_parts.append(p.relative_to(PROJECT).as_posix())
            except (KeyError,ValueError,TypeError,json.JSONDecodeError) as exc:errors.append({'batch':b['batch_id'],'file':str(p),'error':str(exc)})
        if covered==b['count']:accepted.append(b['batch_id'])
        else:
            pending.append(b['batch_id'])
            if covered:partial.append({'batch':b['batch_id'],'completed_rows':covered,'expected_rows':b['count']})
    write_json(PROJECT/'reports/draft-merge.json',{'accepted':accepted,'accepted_parts':accepted_parts,'partial':partial,'pending':pending,'errors':errors})
    if errors:raise SystemExit('Draft validation failed; canonical translations were not written. See reports/draft-merge.json')
    write_jsonl(PROJECT/'translations/zh-CN.jsonl',[records[u['id']] for u in units.values()])
    print('Merged full batches:',len(accepted),'committed small parts:',len(accepted_parts),'pending batches:',len(pending),'draft units:',sum(bool(r.get('draft')) for r in records.values()))


def export_workbook(output=None):
    from openpyxl import Workbook,load_workbook
    from openpyxl.styles import Font,Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter
    units=load_units();records=current_translations();wb=Workbook();wb.remove(wb.active)
    columns=['id','source_sha256','source','initial_zh','review_zh','status','notes','reviewer','usage_approved','model',
             'category','variant','speaker','context_examples','occurrence_count','width_hint_px','review_flags']
    expected={}
    def sheet(name,cols,rows,status_values=None):
        ws=wb.create_sheet(name);ws.append(cols)
        for row in rows:
            ws.append([json.dumps(row.get(k),ensure_ascii=False) if isinstance(row.get(k),(dict,list)) else row.get(k,'') for k in cols])
        for cells in ws:
            for cell in cells:
                if isinstance(cell.value,str):
                    if len(cell.value)>32767:raise ValueError('Excel text length limit: '+str(cell.coordinate))
                    cell.data_type='s'
        for cell in ws[1]:cell.font=Font(bold=True)
        ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions
        for i,col in enumerate(cols,1):
            ws.column_dimensions[get_column_letter(i)].width=60 if col in ('source','initial_zh','review_zh') else 44 if col in ('context_examples','notes','path') else 24
        if 'status' in cols and status_values:
            dv=DataValidation(type='list',formula1='"'+','.join(status_values)+'"',allow_blank=False)
            ws.add_data_validation(dv);letter=get_column_letter(cols.index('status')+1);dv.add(f'{letter}2:{letter}{max(2,ws.max_row)}')
        return ws
    for category in ['ui','vsn','dialogue','code','image']:
        rows=[]
        for u in units:
            if u['category']!=category:continue
            rec=records[u['id']]
            rows.append({**u,'initial_zh':rec['draft'],'review_zh':rec['translation'],'status':rec['status'],
                         'notes':rec['notes'],'reviewer':rec['reviewer'],'usage_approved':rec.get('usage_approved',False),'model':rec.get('model','')})
        if not rows:continue
        sheet(category,columns,rows,STATES);expected[category]=rows
    overrides=read_csv(PROJECT/'review/occurrence_overrides.csv')
    sheet('occurrence_overrides',list(overrides[0]),overrides,STATES)
    fixed=read_csv(PROJECT/'review/fixed_images.csv')
    sheet('fixed_images',list(fixed[0]) if fixed else ['asset_id'],fixed,
          ['needs_visual_review','needs_transcription','needs_redraw','draft','reviewed','approved','not_text'])
    glossary=read_csv(PROJECT/'review/glossary.csv');sheet('glossary',list(glossary[0]),glossary,['proposed','approved','rejected'])
    instructions=[{'item':'主表','value':'initial_zh 为模型初翻，来源见 model 列；在 review_zh 填校对稿，设置 reviewed/approved。勿更改 ID、原文或哈希。'},
                  {'item':'覆盖','value':'已按 patch.pak 覆盖主 JAR；完整位置见 catalog/occurrences.jsonl，不丢弃重复位置。'},
                  {'item':'NSFW','value':'忠实对应原文，不擅自删改或扩写；歧义在 notes 记录。'},
                  {'item':'图片','value':'fixed_images 只收录资产/转录/译文；不会自动重绘图片。浏览 review/image_browser.html。'},
                  {'item':'代码','value':'code 是候选，须确认确实面向玩家；usage_approved=true 才允许写入。'},
                  {'item':'单位置差异','value':'编辑 review/occurrence_overrides.csv，按 occurrence_id 覆盖个别位置。'},
                  {'item':'宽度','value':'width_hint_px 只是校对提示，不是已验证的 UI 容器宽度；不要为过宽自动删减语义。'}]
    sheet('instructions',['item','value'],instructions)
    p=Path(output).resolve() if output else PROJECT/'review/翻译校对.xlsx'
    if not p.resolve().is_relative_to((PROJECT/'review').resolve()):raise ValueError('Workbook output must stay in project/review')
    if output and p.exists():raise ValueError('Refuse to overwrite a named review snapshot')
    p.parent.mkdir(parents=True,exist_ok=True);wb.save(p)
    check=load_workbook(p,data_only=False)
    for cat,rows in expected.items():
        ws=check[cat];assert ws.max_row==len(rows)+1
        for i,row in enumerate(rows,2):assert ws.cell(i,1).value==row['id'] and ws.cell(i,3).value==row['source']
    assert not any(cell.data_type=='f' for ws in check for row in ws for cell in row)
    check.close();print('Workbook saved and reopened:',p)


def import_workbook(path):
    from openpyxl import load_workbook
    units=load_units();by_id={u['id']:u for u in units};records=current_translations();seen=set();metrics=glyph_metrics()
    wb=load_workbook(path,data_only=False)
    if any(c.data_type=='f' for ws in wb for row in ws for c in row):raise ValueError('Formulas are not accepted in review workbooks')
    for category in ['ui','vsn','dialogue','code','image']:
        if category not in wb.sheetnames:continue
        ws=wb[category];cols=[c.value for c in ws[1]]
        for values in ws.iter_rows(min_row=2,values_only=True):
            row=dict(zip(cols,values));uid=row['id']
            if uid not in by_id or uid in seen:raise ValueError('Unknown or duplicate ID '+str(uid))
            seen.add(uid);u=by_id[uid];rec=records[uid]
            if row['source']!=u['source'] or row['source_sha256']!=u['source_sha256']:raise ValueError('Source changed '+uid)
            shown_draft=row.get('initial_zh') or ''
            if shown_draft and shown_draft!=rec['draft']:raise ValueError('Initial draft changed '+uid)
            status=row['status'];translation=row.get('review_zh') or ''
            stale_blank=not shown_draft and bool(rec['draft'])
            if stale_blank and status=='untranslated' and not translation and not row.get('notes') and not row.get('reviewer') and str(row.get('usage_approved')).lower()!='true':
                continue  # Older review snapshots must not erase newly arrived machine drafts.
            if stale_blank and status=='untranslated' and not translation:
                status=rec['status'];translation=rec.get('translation','')
            if status not in STATES:raise ValueError('Unknown review status '+str(status))
            if status in ('reviewed','approved'):
                errors=[q for q in validate_translation(u,translation,metrics) if q[0]=='error']
                if errors:raise ValueError('Review has blocking errors '+uid+' '+str(errors))
            rec.update(translation=translation,status=status,notes=row.get('notes') or (rec.get('notes','') if stale_blank else ''),reviewer=row.get('reviewer') or (rec.get('reviewer','') if stale_blank else ''),
                       usage_approved=str(row.get('usage_approved')).lower()=='true')
    if seen!=set(by_id):raise ValueError('Workbook is missing unit rows')
    fixed=read_csv(PROJECT/'review/fixed_images.csv');fixed_by_id={r['asset_id']:r for r in fixed};fixed_seen=set()
    if 'fixed_images' not in wb:raise ValueError('Missing fixed_images sheet')
    ws=wb['fixed_images'];cols=[c.value for c in ws[1]]
    editable={'source_transcription','translation','status','review_notes'}
    for values in ws.iter_rows(min_row=2,values_only=True):
        row=dict(zip(cols,values));aid=row['asset_id']
        if aid not in fixed_by_id or aid in fixed_seen:raise ValueError('Unknown/duplicate image asset ID')
        fixed_seen.add(aid);old=fixed_by_id[aid]
        for key in cols:
            if key not in editable and str(row.get(key) or '')!=str(old.get(key) or ''):raise ValueError('Image mapping changed '+aid+' '+key)
        for key in editable:old[key]=row.get(key) or ''
    if fixed_seen!=set(fixed_by_id):raise ValueError('Missing image rows')
    glossary=read_csv(PROJECT/'review/glossary.csv');gloss_by_term={r['term']:r for r in glossary};term_seen=set()
    if 'glossary' not in wb:raise ValueError('Missing glossary sheet')
    ws=wb['glossary'];cols=[c.value for c in ws[1]]
    for values in ws.iter_rows(min_row=2,values_only=True):
        row=dict(zip(cols,values));term=row['term']
        if term not in gloss_by_term or term in term_seen:raise ValueError('Unknown/duplicate glossary term')
        term_seen.add(term)
        if row.get('status') not in ('proposed','approved','rejected'):raise ValueError('Invalid glossary status')
        gloss_by_term[term].update({k:row.get(k) or '' for k in ('usage','proposed_zh','note','status')})
    if term_seen!=set(gloss_by_term):raise ValueError('Missing glossary terms')
    overrides=read_csv(PROJECT/'review/occurrence_overrides.csv');override_by_id={r['occurrence_id']:r for r in overrides};override_seen=set()
    if 'occurrence_overrides' not in wb:raise ValueError('Missing occurrence_overrides sheet')
    ws=wb['occurrence_overrides'];cols=[c.value for c in ws[1]]
    for values in ws.iter_rows(min_row=2,values_only=True):
        row=dict(zip(cols,values));oid=row['occurrence_id']
        if oid not in override_by_id or oid in override_seen:raise ValueError('Unknown/duplicate occurrence ID')
        override_seen.add(oid);old=override_by_id[oid]
        for key in cols:
            if key not in ('translation_override','status','notes') and str(row.get(key) or '')!=str(old.get(key) or ''):raise ValueError('Occurrence mapping changed '+oid)
        status=row.get('status');text=row.get('translation_override') or ''
        if status not in STATES:raise ValueError('Invalid occurrence status')
        if status in ('reviewed','approved') and text:
            errors=[q for q in validate_translation(by_id[old['unit_id']],text,metrics) if q[0]=='error']
            if errors:raise ValueError('Invalid occurrence translation '+oid+' '+str(errors))
        old.update(translation_override=text,status=status,notes=row.get('notes') or '')
    if override_seen!=set(override_by_id):raise ValueError('Missing occurrence rows')
    # All identifiers/source guards pass before any canonical file is changed.
    write_jsonl(PROJECT/'translations/zh-CN.jsonl',[records[u['id']] for u in units])
    write_csv(PROJECT/'review/occurrence_overrides.csv',overrides,list(overrides[0]))
    write_csv(PROJECT/'review/fixed_images.csv',fixed,list(fixed[0]))
    write_csv(PROJECT/'review/glossary.csv',glossary,list(glossary[0]));wb.close()
    print('Imported',len(seen),'review rows,',len(fixed),'image rows and',len(glossary),'terms')


def validate():
    units=load_units();records=current_translations();metrics=glyph_metrics();issues=[];counts=collections.Counter()
    for u in units:
        rec=records[u['id']]
        if rec['source_sha256']!=u['source_sha256']:raise ValueError('Source identity mismatch')
        text=rec['translation'] or rec['draft']
        if not text:
            counts['blocked' if rec['status']=='blocked' else 'untranslated']+=1
            if rec['status']=='blocked':issues.append({'id':u['id'],'category':u['category'],'severity':'warning','kind':'blocked','detail':rec.get('notes',''),'source':u['source'],'translation':''})
            continue
        counts[rec['status']]+=1
        for severity,kind,detail in validate_translation(u,text,metrics):issues.append({'id':u['id'],'category':u['category'],'severity':severity,'kind':kind,'detail':detail,'source':u['source'],'translation':text})
    write_csv(PROJECT/'reports/translation-qa.csv',issues,['id','category','severity','kind','detail','source','translation'])
    summary={'counts':dict(counts),'issues':dict(collections.Counter(r['severity']+':'+r['kind'] for r in issues)),
             'translated_units':sum(bool(r['draft'] or r['translation']) for r in records.values()),'total_units':len(units)}
    write_json(PROJECT/'reports/translation-qa-summary.json',summary);print(json.dumps(summary,ensure_ascii=False,indent=2));return issues


def selected_translations(include_drafts=False):
    units={u['id']:u for u in load_units()};records=current_translations();metrics=glyph_metrics();out={}
    for uid,u in units.items():
        text=merged_value(records[uid],include_drafts)
        if text is None:continue
        errors=[q for q in validate_translation(u,text,metrics) if q[0]=='error']
        if errors:raise ValueError('Translation blocked by QA '+uid+' '+str(errors))
        out[uid]=text
    return out


def render_resources(translations,overrides=None):
    """Return logical path -> bytes only for effective text assets; no source mutation."""
    units={u['id']:u for u in load_units()};resources={r['logical_path']:r for r in read_jsonl(PROJECT/'catalog/resources.jsonl')}
    by_path=collections.defaultdict(list)
    for o in load_occurrences():
        if o['category'] in ('dialogue','vsn','ui'):by_path[o['logical_path']].append(o)
    output={}
    for path,occurrences in by_path.items():
        original=resource_path(resources[path]).read_bytes();changed=False
        if path.endswith('.json'):
            data=json.loads(original)
            for o in occurrences:
                raw=pointer_get(data,o['json_pointer']) if 'json_pointer' in o else data[o['json_key']];u=units[o['unit_id']]
                assert source_hash(raw)==o['raw_payload_sha256']
                text=(overrides or {}).get(o['occurrence_id'],translations.get(u['id']))
                if text is not None:
                    value=o['leading_ws']+text+o['trailing_ws']
                    if 'json_pointer' in o:pointer_set(data,o['json_pointer'],value)
                    else:data[o['json_key']]=value
                    changed=True
            result=json.dumps(data,ensure_ascii=False,separators=(',',':')).encode('utf8') if changed else original
        else:
            lines=original.decode('utf8').splitlines(keepends=True)
            for o in occurrences:
                u=units[o['unit_id']];i=o['line']-1;line=lines[i]
                eol='\r\n' if line.endswith('\r\n') else '\n' if line.endswith('\n') else ''
                body=line[:-len(eol)] if eol else line
                assert body==o['prefix']+o['leading_ws']+u['source']+o['trailing_ws'],o['location']
                text=(overrides or {}).get(o['occurrence_id'],translations.get(u['id']))
                if text is not None:lines[i]=o['prefix']+o['leading_ws']+text+o['trailing_ws']+eol;changed=True
            result=''.join(lines).encode('utf8') if changed else original
        output[path]=result
    return output


def render(path,include_drafts=False):
    dest=Path(path).resolve();allowed=(PROJECT/'builds').resolve()
    if not dest.is_relative_to(allowed) or dest==allowed:raise ValueError('Output must be a new child of project/builds')
    if dest.exists():raise ValueError('Refuse to overwrite existing output')
    translations=selected_translations(include_drafts);override_map={}
    occurrences={o['occurrence_id']:o for o in load_occurrences()};units={u['id']:u for u in load_units()}
    for row in read_csv(PROJECT/'review/occurrence_overrides.csv'):
        if row['status'] in ('reviewed','approved') and row['translation_override']:
            o=occurrences[row['occurrence_id']];errors=[q for q in validate_translation(units[o['unit_id']],row['translation_override'],glyph_metrics()) if q[0]=='error']
            if errors:raise ValueError('Invalid occurrence override '+str(errors))
            override_map[o['occurrence_id']]=row['translation_override']
    dest.mkdir(parents=True);manifest=[];resources={r['logical_path']:r for r in read_jsonl(PROJECT/'catalog/resources.jsonl')}
    for key,data in render_resources(translations,override_map).items():
        r=resources[key];name=r['resource_id']+Path(key).suffix;(dest/name).write_bytes(data)
        manifest.append({'logical_path':key,'archive':r['archive'],'archive_entry':r['archive_entry'],'file':name,'sha256':digest(data),
                         'changed':digest(data)!=r['sha256']})
    write_json(dest/'manifest.json',manifest);print('Rendered resources:',len(manifest),'changed:',sum(r['changed'] for r in manifest))


def main():
    ap=argparse.ArgumentParser();sub=ap.add_subparsers(dest='action',required=True)
    for action in ['prepare','merge-drafts','validate']:sub.add_parser(action)
    p=sub.add_parser('export');p.add_argument('--output')
    p=sub.add_parser('import-review');p.add_argument('workbook')
    p=sub.add_parser('render');p.add_argument('output');p.add_argument('--include-drafts',action='store_true')
    a=ap.parse_args()
    if a.action=='prepare':prepare()
    elif a.action=='export':export_workbook(a.output)
    elif a.action=='merge-drafts':merge_drafts()
    elif a.action=='validate':validate()
    elif a.action=='import-review':import_workbook(a.workbook)
    elif a.action=='render':render(a.output,a.include_drafts)


if __name__=='__main__':sys.stdout.reconfigure(encoding='utf8');main()
