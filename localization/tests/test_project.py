"""Deterministic project checks; no game launch, network or install writes."""
from pathlib import Path
import collections
import json
import sys
import unittest
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'tools'))
from projectlib import *
from review import render_resources,import_workbook,current_translations
from PIL import Image


class ProjectTests(unittest.TestCase):
    def test_identity_and_complete_locations(self):
        units=load_units();occ=load_occurrences();ids={u['id'] for u in units}
        self.assertEqual(len(ids),len(units));self.assertEqual(len({o['occurrence_id'] for o in occ}),len(occ))
        self.assertTrue(all(o['unit_id'] in ids for o in occ))
        counts=collections.Counter(o['unit_id'] for o in occ)
        for u in units:
            self.assertEqual(u['source_sha256'],source_hash(u['source']))
            self.assertEqual(u['occurrence_count'],counts[u['id']])
        self.assertEqual(read_json(PROJECT/'reports/unclassified_text_lines.json'),[])

    def test_source_snapshots_and_archive_hashes(self):
        checked=set()
        for r in read_jsonl(PROJECT/'catalog/resources.jsonl'):
            for v in r['versions']:
                if v['materialized'] and v['sha256'] not in checked:
                    self.assertEqual(digest((PROJECT/v['blob_path']).read_bytes()),v['sha256']);checked.add(v['sha256'])
            self.assertEqual(r['archive'],r['versions'][-1]['archive'])
        for a in read_json(PROJECT/'sources/archives.json'):
            self.assertEqual(digest(Path(a['path']).read_bytes()),a['sha256'])

    def test_byte_exact_no_translation_roundtrip(self):
        resources={r['logical_path']:r for r in read_jsonl(PROJECT/'catalog/resources.jsonl')}
        rendered=render_resources({})
        self.assertEqual(len(rendered),71)
        for key,data in rendered.items():self.assertEqual(data,resource_path(resources[key]).read_bytes(),key)

    def test_newline_and_prefix_preservation(self):
        units={u['id']:u for u in load_units()}
        o=next(o for o in load_occurrences() if o['category']=='vsn' and o['logical_path']=='vsn/ClassicPreTitle/script.vsn')
        text=render_resources({o['unit_id']:'中文测试。'})[o['logical_path']].decode('utf8')
        line=text.splitlines()[o['line']-1]
        self.assertEqual(line,o['prefix']+o['leading_ws']+'中文测试。'+o['trailing_ws'])
        original=next(r for r in read_jsonl(PROJECT/'catalog/resources.jsonl') if r['logical_path']==o['logical_path'])
        self.assertEqual(text.count('\r\n'),resource_path(original).read_bytes().count(b'\r\n'))

    def test_review_book_noop_import_and_source_guard(self):
        from openpyxl import load_workbook
        p=PROJECT/'review/翻译校对.xlsx';before=current_translations()
        import_workbook(p);self.assertEqual(before,current_translations())
        wb=load_workbook(p,data_only=False)
        self.assertFalse(any(c.data_type=='f' for ws in wb for row in ws for c in row))
        wb['ui']['C2']=str(wb['ui']['C2'].value)+' SOURCE TAMPER'
        dest=PROJECT/'tests/output';dest.mkdir(exist_ok=True)
        bad=dest/'tampered-source.xlsx';wb.save(bad);wb.close()
        with self.assertRaises(ValueError):import_workbook(bad)
        self.assertEqual(before,current_translations())

    def test_older_workbook_keeps_newly_arrived_draft(self):
        import copy
        from unittest.mock import patch
        import review
        before=current_translations();future=copy.deepcopy(before)
        uid=next(uid for uid,r in future.items() if r['status']=='untranslated')
        future[uid].update(draft='测试初稿',status='draft',notes='Synthetic future-arrival fixture, not a product translation')
        captured={}
        def capture(path,rows):captured.update({r['id']:dict(r) for r in rows})
        with patch.object(review,'current_translations',return_value=future),patch.object(review,'write_jsonl',side_effect=capture),patch.object(review,'write_csv'):
            review.import_workbook(PROJECT/'review/翻译校对.xlsx')
        self.assertEqual(captured[uid]['draft'],'测试初稿');self.assertEqual(captured[uid]['status'],'draft')
        self.assertEqual(before,current_translations())

    def test_literal_equals_never_becomes_formula(self):
        from openpyxl import Workbook,load_workbook
        wb=Workbook();ws=wb.active;ws['A1']='=1+1';ws['A1'].data_type='s'
        dest=PROJECT/'tests/output';dest.mkdir(exist_ok=True);p=dest/'literal-equals.xlsx';wb.save(p)
        other=load_workbook(p,data_only=False);self.assertEqual(other.active['A1'].value,'=1+1');self.assertEqual(other.active['A1'].data_type,'s');other.close()

    def test_image_pixels_and_animation_mapping(self):
        resources={r['logical_path']:r for r in read_jsonl(PROJECT/'catalog/resources.jsonl')}
        page=Image.open(resource_path(resources['dataq/sprites.png'])).convert('RGBA')
        fixed=read_jsonl(PROJECT/'catalog/fixed_images.jsonl')
        for r in fixed:
            if r['name'] not in ('ClassicPretitleGamepad','CardViewMenuBGTitle','CosmoTitleData'):continue
            x,y=r['xy'];w,h=r['size'];self.assertEqual(r['rotation'],0)
            with Image.open(PROJECT/r['path']) as crop:self.assertEqual(crop.convert('RGBA').tobytes(),page.crop((x,y,x+w,y+h)).tobytes())
        for a in read_jsonl(PROJECT/'catalog/scene_animations.jsonl'):
            self.assertEqual(a['frame_count'],len(a['frames']))
            self.assertIn('first_frame_preview',a)
            self.assertTrue((PROJECT/a['first_frame_preview']).is_file())
        page.close()

    def test_nested_script_pack_text_and_json_pointer(self):
        resources={r['logical_path']:r for r in read_jsonl(PROJECT/'catalog/resources.jsonl')}
        key='datay/scripts_pack.json';original=json.loads(resource_path(resources[key]).read_bytes())
        o=next(o for o in load_occurrences() if o.get('logical_path')==key and '/shops/male/' in o.get('json_pointer',''))
        changed=json.loads(render_resources({o['unit_id']:'测试译文'})[key])
        self.assertEqual(pointer_get(changed,o['json_pointer']),o['leading_ws']+'测试译文'+o['trailing_ws'])
        self.assertEqual(changed['enemyData'],original['enemyData'])
        self.assertEqual(changed['shops']['female'],original['shops']['female'])
        self.assertEqual(changed['patchDictionary'],original['patchDictionary'])
        self.assertEqual(sum('runtime_overlay_pointer' in x for x in load_occurrences()),107)

    def test_small_part_input_identity(self):
        from draft_part import make_part
        units={u['id']:u for u in load_units()};b=next(b for b in read_json(PROJECT/'batches/manifest.json')['batches'] if b['batch_id']=='batch_005')
        path,count=make_part(b,0,units);rows=read_jsonl(path)
        self.assertEqual(count,40);self.assertEqual([r['id'] for r in rows],b['ids'][:40])
        for row in rows:self.assertEqual(row['source_sha256'],source_hash(row['source']))
        with self.assertRaises(ValueError):make_part(b,1,units)

    def test_draft_model_provenance_does_not_default_to_grok(self):
        registry={'drafts/test-grok.jsonl':{'provider':'xai-oauth','model':'grok-4.7','reasoning_effort':'high'},
                  'drafts/test-mimo.jsonl':{'provider':'xiaomi-token-plan-cn','model':'mimo-v2.6-pro','reasoning_effort':'high'},
                  'drafts/test-sol.jsonl':{'provider':'openai-codex','model':'gpt-6.1-sol','reasoning_effort':'medium'}}
        self.assertEqual(draft_model_label(PROJECT/'drafts/test-grok.jsonl',registry),'xai-oauth/grok-4.7/high')
        self.assertEqual(draft_model_label(PROJECT/'drafts/test-mimo.jsonl',registry),'xiaomi-token-plan-cn/mimo-v2.6-pro/high')
        self.assertEqual(draft_model_label(PROJECT/'drafts/test-sol.jsonl',registry),'openai-codex/gpt-6.1-sol/medium')
        with self.assertRaises(ValueError):draft_model_label(PROJECT/'drafts/unregistered.jsonl',registry)

    def test_marker_and_glyph_checks(self):
        metrics=glyph_metrics();unit={'source':'[TEST] %s|Done','width_hint_px':240}
        self.assertFalse(any(q[0]=='error' for q in validate_translation(unit,'[测试] %s|完成',metrics)))
        bad=validate_translation(unit,'测试|完成',metrics)
        self.assertTrue(any(q[1]=='marker_placeholders' for q in bad))
        self.assertTrue(any(q[1]=='marker_open_square' for q in bad))
        self.assertIn('@EmuButtons',protected('Press @EmuButtons now'))
        self.assertTrue(any(q[1]=='marker_underscore_runs' for q in validate_translation({'source':'_____','width_hint_px':240},'___',metrics)))
        self.assertTrue(any(q[1]=='missing_glyph' for q in validate_translation({'source':'Hello','width_hint_px':240},'你好🙂',metrics)))


if __name__=='__main__':
    sys.stdout.reconfigure(encoding='utf8');unittest.main(verbosity=2)
